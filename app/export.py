"""棚卸結果を紙の管理表に近いレイアウトの .xlsx にする。"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import logic

THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
EXPIRED_FILL = PatternFill("solid", fgColor="F8CBAD")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)
STATUS_TEXT = {"ok": "レ", "warn": "×（1ヶ月未満）", "expired": "×（期限切れ）", "none": ""}
MODE_NOTE = {"none": "期限指定無し"}


def jdate(s: str | None) -> str:
    if not s:
        return ""
    d = date.fromisoformat(s)
    return f"{d.year}/{d.month}/{d.day}"


def counted_range(items) -> str:
    days = sorted({x["counted_on"] for x in items if x["counted_on"]})
    if not days:
        return "未実施"
    return jdate(days[0]) if len(days) == 1 else f"{jdate(days[0])}〜{jdate(days[-1])}"


def safe_title(name: str, used: set) -> str:
    t = "".join(ch for ch in name if ch not in '[]:*?/\\')[:28] or "分類"
    base, n = t, 2
    while t in used:
        t = f"{base[:25]}_{n}"
        n += 1
    used.add(t)
    return t


def cell(ws, r, c, v, **style):
    x = ws.cell(row=r, column=c, value=v)
    for k, val in style.items():
        setattr(x, k, val)
    return x


def build(con, ym: str, month_items, roles: list[str]) -> bytes:
    y, m = map(int, ym.split("-"))
    me, limit = logic.month_end(ym), logic.ok_limit(ym)
    rule = (f"賞味期限管理：確認月の月末（{me.month}/{me.day}）から{logic.WARN_MONTHS}ヶ月"
            f"（{limit.year}/{limit.month}/{limit.day}）以上あるか確認　・{logic.WARN_MONTHS}ヶ月以上あり：レ　・未満：×")
    rows = month_items(con, ym)
    cats = [dict(r) for r in con.execute("SELECT * FROM categories WHERE active=1 ORDER BY sort")]
    appr = {}
    for a in con.execute("SELECT * FROM approvals WHERE ym=?", (ym,)):
        appr[(a["category_id"], a["role"])] = a

    wb = Workbook()
    used: set = set()

    # 1枚目：期限警告の一覧
    ws = wb.active
    ws.title = safe_title("期限警告一覧", used)
    cell(ws, 1, 1, f"期限警告一覧（{y}年{m}月 棚卸）", font=Font(size=14, bold=True))
    cell(ws, 2, 1, rule)
    heads = ["分類", "コード", "品名", "総重量(kg)", "区分", "期限", "判定", "対応", "対応予定日", "対応メモ", "担当者", "棚卸日"]
    for i, h in enumerate(heads, 1):
        cell(ws, 4, i, h, fill=HEAD_FILL, border=BORDER, alignment=CENTER, font=Font(bold=True))
    r = 5
    for it in sorted((x for x in rows if x["status"] in ("warn", "expired")), key=lambda x: x["expiry_date"] or ""):
        vals = [it["category_name"], it["code"], it["name"], it["total_kg"], it["expiry_kind"],
                jdate(it["expiry_date"]), STATUS_TEXT[it["status"]], it["action"], jdate(it["action_date"]),
                it["action_note"], it["counted_by"], jdate(it["counted_on"])]
        fill = EXPIRED_FILL if it["status"] == "expired" else WARN_FILL
        for i, v in enumerate(vals, 1):
            cell(ws, r, i, v, border=BORDER, alignment=WRAP, fill=fill)
        r += 1
    if r == 5:
        cell(ws, 5, 1, "警告はありません")
    for i, w in enumerate([16, 10, 32, 11, 6, 12, 15, 16, 12, 28, 10, 11], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 分類ごとのシート
    for cat in cats:
        items = [x for x in rows if x["category_id"] == cat["id"]]
        if not items:
            continue
        split = cat["locations"] == "split"
        ws = wb.create_sheet(safe_title(cat["name"], used))
        cell(ws, 1, 1, f"在庫賞味期限管理（{cat['name']}）", font=Font(size=14, bold=True))
        cell(ws, 2, 1, f"棚卸月：{y}年{m}月　　棚卸日：{counted_range(items)}　　出力日：{date.today():%Y/%m/%d}")
        cell(ws, 3, 1, rule)

        # 確認欄（右上）
        start = 13 if split else 10
        for i, role in enumerate(roles):
            a = appr.get((cat["id"], role))
            cell(ws, 1, start + i, role, fill=HEAD_FILL, border=BORDER, alignment=CENTER, font=Font(bold=True))
            txt = f"{a['staff_name']}\n{a['at'][:10]}" + (f"\n{a['detail']}" if a and a["detail"] else "") if a else ""
            cell(ws, 2, start + i, txt, border=BORDER, alignment=CENTER)
        ws.row_dimensions[2].height = 45

        loc_heads = (["資材室", "", "倉庫/パレット", "", ""] if split else ["在庫", ""])
        heads = ["コード", "品名", "保管", "ケース重量(kg)", *loc_heads, "総重量(kg)", "区分",
                 "賞味期限/使用期限", "チェック(レ・×)", "対応", "対応メモ", "担当者", "棚卸日"]
        loc_sub = ["ケース", "端数kg", "パレット", "ケース", "端数kg"] if split else ["ケース", "端数kg"]
        sub = ["", "", "", "", *loc_sub, "", "", "", "", "", "", "", ""]
        for i, (h, s) in enumerate(zip(heads, sub), 1):
            cell(ws, 5, i, h, fill=HEAD_FILL, border=BORDER, alignment=CENTER, font=Font(bold=True))
            cell(ws, 6, i, s, fill=HEAD_FILL, border=BORDER, alignment=CENTER, font=Font(bold=True))
        ws.merge_cells(start_row=5, start_column=5, end_row=5, end_column=6)
        if split:
            ws.merge_cells(start_row=5, start_column=7, end_row=5, end_column=9)
        for i, h in enumerate(heads, 1):
            if h and not sub[i - 1]:
                ws.merge_cells(start_row=5, start_column=i, end_row=6, end_column=i)

        r = 7
        for it in items:
            entered = it["entry_id"] is not None
            loc = ([it["room_cases"], it["room_kg"], it["wh_pallets"], it["wh_cases"], it["wh_kg"]] if split
                   else [it["room_cases"], it["room_kg"]]) if entered else [None] * len(loc_sub)
            expiry = jdate(it["expiry_date"]) or (MODE_NOTE.get(it["expiry_mode"], "") if entered else "")
            if it["mfg_date"]:
                expiry += f"\n(製造日{jdate(it['mfg_date'])})"
            name = it["name"] + ("（終売）" if not it["active"] else "")
            vals = [it["code"], name, it["storage"], it["entry_case_weight"] if entered else it["case_weight"], *loc,
                    it["total_kg"] if entered else None, it["expiry_kind"] if entered else "", expiry,
                    STATUS_TEXT[it["status"]] if entered else "未入力", it["action"] or "", it["action_note"] or "",
                    it["counted_by"] or "", jdate(it["counted_on"])]
            fill = {"warn": WARN_FILL, "expired": EXPIRED_FILL}.get(it["status"]) if entered else None
            for i, v in enumerate(vals, 1):
                x = cell(ws, r, i, v, border=BORDER, alignment=WRAP if i in (2, 15 if split else 12) else CENTER)
                if fill:
                    x.fill = fill
            r += 1

        widths = [10, 30, 6, 9, *([7, 8, 7, 7, 8] if split else [7, 8]), 10, 5, 13, 11, 14, 22, 9, 11]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "C7"
        ws.print_title_rows = "5:6"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
