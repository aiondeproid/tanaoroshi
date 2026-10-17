"""予想在庫のExcel（ひな形の出力と取り込み）。"""
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from .export import BORDER, CENTER, HEAD_FILL, cell

HEADERS = ["コード", "品名", "分類", "保管", "予想在庫(kg)", "時点"]
WIDTHS = [12, 36, 18, 8, 14, 12]


def norm(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(v))).strip()


def norm_code(v) -> str:
    """tools/import_master.py と同じ正規化。Excelで数値になって消えた先頭の0も戻す。"""
    c = norm(v).upper()
    return c.zfill(5) if c.isdigit() and len(c) < 5 else c


def build_template(ym: str, rows: list[dict]) -> bytes:
    y, m = map(int, ym.split("-"))
    wb = Workbook()
    ws = wb.active
    ws.title = "予想在庫"
    cell(ws, 1, 1, f"{y}年{m}月 予想在庫")
    cell(ws, 2, 1, "「予想在庫(kg)」と、いつの時点の数字かを「時点」に入れて、管理画面の「予想在庫」タブから取り込んでください。"
                   "予想在庫が空欄の行は変更しません。時点が空欄の行は取り込み画面で選んだ日になります。")
    for c, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell(ws, 3, c, h, fill=HEAD_FILL, border=BORDER, alignment=CENTER)
        ws.column_dimensions[chr(64 + c)].width = w
    for r, it in enumerate(rows, 4):
        as_of = date.fromisoformat(it["expected_as_of"]) if it.get("expected_as_of") else None
        vals = [it["code"], it["name"], it["category_name"], it["storage"], it["expected_kg"], as_of]
        for c, v in enumerate(vals, 1):
            cell(ws, r, c, v, border=BORDER)
        ws.cell(row=r, column=1).number_format = "@"  # コードの先頭の0が消えないよう文字列にする
        ws.cell(row=r, column=5).number_format = "0.###"
        ws.cell(row=r, column=6).number_format = "yyyy/m/d"
    ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_as_of(v) -> str | None:
    """時点のセルを YYYY-MM-DD にする。空欄は None、読めなければ ValueError。"""
    if v is None or norm(v) == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    m = re.fullmatch(r"(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})日?(?:時点)?", norm(v).replace(" ", ""))
    if not m:
        raise ValueError(v)
    return date(*map(int, m.groups())).isoformat()


def parse(data: bytes) -> tuple[dict[str, tuple[int, float, str | None]], list[str]]:
    """{コード: (行番号, kg, 時点)} とエラーの一覧を返す。予想在庫が空欄の行は無視する。
    時点の列は無くてもよく、空欄なら None。"""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except (zipfile.BadZipFile, KeyError, OSError):
        return {}, ["Excelファイル（.xlsx）を選んでください。古い .xls は、Excelで .xlsx に保存し直してください"]
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    code_col = kg_col = as_of_col = header = None
    for i, row in enumerate(rows[:20]):
        texts = [norm(v) if v is not None else "" for v in row]
        codes = [c for c, t in enumerate(texts) if t == "コード"]
        kgs = [c for c, t in enumerate(texts) if "予想" in t]
        if codes and kgs:
            code_col, kg_col, header = codes[0], kgs[0], i
            as_of_col = next((c for c, t in enumerate(texts) if t in ("時点", "基準日")), None)
            break
    if header is None:
        return {}, ["見出しの行が見つかりません。「コード」と「予想在庫(kg)」の列がある表にしてください"]

    out: dict[str, tuple[int, float, str | None]] = {}
    errors: list[str] = []
    for i, row in enumerate(rows[header + 1:], header + 2):
        get = lambda c: row[c] if c < len(row) else None  # noqa: E731
        raw_code, raw_kg = get(code_col), get(kg_col)
        if raw_kg is None or norm(raw_kg) == "":
            continue
        code = norm_code(raw_code) if raw_code is not None else ""
        if not code:
            errors.append(f"{i}行目：コードが空です")
            continue
        if isinstance(raw_kg, (int, float)):
            kg = float(raw_kg)
        else:
            try:
                kg = float(norm(raw_kg).replace(",", "").removesuffix("kg").removesuffix("KG").strip())
            except ValueError:
                errors.append(f"{i}行目（{code}）：予想在庫「{raw_kg}」が数字ではありません")
                continue
        if kg < 0:
            errors.append(f"{i}行目（{code}）：予想在庫がマイナスです")
            continue
        try:
            as_of = parse_as_of(get(as_of_col)) if as_of_col is not None else None
        except ValueError:
            errors.append(f"{i}行目（{code}）：時点「{get(as_of_col)}」が日付ではありません")
            continue
        if code in out:
            errors.append(f"{i}行目（{code}）：{out[code][0]}行目と同じコードです")
            continue
        out[code] = (i, round(kg, 3), as_of)
    return out, errors
