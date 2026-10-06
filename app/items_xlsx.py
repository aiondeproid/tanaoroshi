"""品目マスタのExcel（一覧の出力と、まとめて変更の取り込み）。"""
import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .expected_xlsx import norm, norm_code
from .export import BORDER, CENTER, HEAD_FILL, cell

EXPIRY_MODE = {"date": "期限を入力", "mfg": "製造日から計算", "none": "期限指定無し"}
ACTIVE = {1: "有効", 0: "終売"}

# (キー, 見出し, 幅)。見出しは取り込み時に列を探す目印にもなる
COLUMNS = [
    ("code", "コード", 10),
    ("name", "品名", 34),
    ("category", "分類", 22),
    ("storage", "保管", 8),
    ("case_weight", "ケース重量(kg)", 12),
    ("pallet_cases", "1パレットのケース数", 12),
    ("expiry_mode", "期限の種類", 15),
    ("mfg_months", "製造日からの月数", 10),
    ("note", "メモ", 28),
    ("sort", "並び順", 8),
    ("active", "状態", 8),
]
# 空欄にすると「未設定」になる列。それ以外の列は空欄なら今の値のまま
NULLABLE = {"case_weight", "pallet_cases", "mfg_months", "note"}


def category_label(c: dict) -> str:
    return f"{c['grp']} / {c['name']}"


def build(items: list[dict], categories: list[dict], storages: list[str]) -> bytes:
    cat_label = {c["id"]: category_label(c) for c in categories}
    wb = Workbook()
    ws = wb.active
    ws.title = "品目"
    cell(ws, 1, 1, "品目マスタ")
    cell(ws, 2, 1, "直したい所を書き換えて、管理画面の「品目」タブから取り込んでください。"
                   "コードで品目を探します（コードは変えないでください）。新しいコードの行は品目の追加になります。"
                   "ケース重量・パレット・月数・メモは空欄にすると未設定になります。行を消しても品目は消えません（終売は「状態」で）。")
    for c, (_, head, width) in enumerate(COLUMNS, 1):
        cell(ws, 3, c, head, fill=HEAD_FILL, border=BORDER, alignment=CENTER)
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, it in enumerate(items, 4):
        vals = {
            **it, "category": cat_label.get(it["category_id"], ""),
            "expiry_mode": EXPIRY_MODE[it["expiry_mode"]], "active": ACTIVE[1 if it["active"] else 0],
        }
        for c, (key, _, _) in enumerate(COLUMNS, 1):
            v = vals[key]
            cell(ws, r, c, v if v != "" else None, border=BORDER)
        ws.cell(row=r, column=1).number_format = "@"  # コードの先頭の0が消えないよう文字列にする
    last = max(len(items) + 3, 3) + 500  # 追加の行にも選択肢が出るよう少し先まで

    # 選択肢は別シートに置いてプルダウンにする
    ls = wb.create_sheet("選択肢")
    lists = {
        "category": [category_label(c) for c in categories],
        "storage": storages,
        "expiry_mode": list(EXPIRY_MODE.values()),
        "active": list(ACTIVE.values()),
    }
    keys = [k for k, _, _ in COLUMNS]
    for lc, (key, values) in enumerate(lists.items(), 1):
        cell(ls, 1, lc, dict((k, h) for k, h, _ in COLUMNS)[key], fill=HEAD_FILL, border=BORDER)
        for lr, v in enumerate(values, 2):
            ls.cell(row=lr, column=lc, value=v)
        ls.column_dimensions[get_column_letter(lc)].width = 26
        col = get_column_letter(keys.index(key) + 1)
        ref = f"'選択肢'!${get_column_letter(lc)}$2:${get_column_letter(lc)}${max(len(values), 1) + 1}"
        dv = DataValidation(type="list", formula1=ref, allow_blank=True)
        dv.add(f"{col}4:{col}{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(COLUMNS))}{max(len(items) + 3, 4)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _number(raw):
    if isinstance(raw, (int, float)):
        return float(raw)
    t = norm(raw).replace(",", "")
    for u in ("kg", "KG", "ケース", "ヶ月", "か月", "カ月"):
        t = t.removesuffix(u).strip()
    return float(t)


def parse(data: bytes, categories: list[dict], storages: list[str]) -> tuple[list[dict], list[str]]:
    """[{line, code, values}] とエラーの一覧を返す。values に無いキーは「今の値のまま」。"""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except (zipfile.BadZipFile, KeyError, OSError):
        return [], ["Excelファイル（.xlsx）を選んでください。古い .xls は、Excelで .xlsx に保存し直してください"]
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))

    heads = {norm(h): k for k, h, _ in COLUMNS}
    col_of: dict[str, int] = {}
    header = None
    for i, row in enumerate(rows[:20]):
        found = {heads[norm(v)]: c for c, v in enumerate(row) if v is not None and norm(v) in heads}
        if "code" in found and len(found) >= 2:
            col_of, header = found, i
            break
    if header is None:
        return [], ["見出しの行が見つかりません。「コード」と、変えたい列（「品名」「ケース重量(kg)」など）の見出しがある表にしてください"]

    by_label = {category_label(c): c["id"] for c in categories}
    names: dict[str, list[int]] = {}
    for c in categories:
        names.setdefault(norm(c["name"]), []).append(c["id"])
    mode_of = {v: k for k, v in EXPIRY_MODE.items()} | {k: k for k in EXPIRY_MODE}
    active_of = {"有効": True, "終売": False}

    out: list[dict] = []
    errors: list[str] = []
    seen: dict[str, int] = {}
    for line, row in enumerate(rows[header + 1:], header + 2):
        get = lambda k: row[col_of[k]] if col_of[k] < len(row) else None  # noqa: E731
        raw = {k: get(k) for k in col_of}
        blank = {k for k, v in raw.items() if v is None or norm(v) == ""}
        if len(blank) == len(raw):
            continue
        if "code" in blank:
            errors.append(f"{line}行目：コードが空です")
            continue
        code = norm_code(raw["code"])
        if code in seen:
            errors.append(f"{line}行目（{code}）：{seen[code]}行目と同じコードです")
            continue
        seen[code] = line
        err = lambda msg: errors.append(f"{line}行目（{code}）：{msg}")  # noqa: E731
        values: dict = {}
        for k, v in raw.items():
            if k == "code":
                continue
            if k in blank:
                if k in NULLABLE:
                    values[k] = "" if k == "note" else None
                continue
            t = norm(v)
            if k == "name":
                values[k] = t
            elif k == "category":
                cid = by_label.get(t) or (names[t][0] if len(names.get(t, [])) == 1 else None)
                if cid is None:
                    err(f"分類「{t}」がありません（「グループ / 分類名」で書いてください）")
                else:
                    values["category_id"] = cid
            elif k == "storage":
                if t not in storages:
                    err(f"保管「{t}」は {'・'.join(storages)} のどれかにしてください")
                else:
                    values[k] = t
            elif k == "expiry_mode":
                if t not in mode_of:
                    err(f"期限の種類「{t}」は {'・'.join(EXPIRY_MODE.values())} のどれかにしてください")
                else:
                    values[k] = mode_of[t]
            elif k == "active":
                if t not in active_of:
                    err(f"状態「{t}」は 有効・終売 のどちらかにしてください")
                else:
                    values[k] = active_of[t]
            elif k == "note":
                values[k] = t
            else:  # 数字の列
                label = dict((kk, h) for kk, h, _ in COLUMNS)[k]
                try:
                    n = _number(v)
                except ValueError:
                    err(f"{label}「{v}」が数字ではありません")
                    continue
                if k in ("mfg_months", "sort"):
                    if not n.is_integer():
                        err(f"{label}は整数にしてください")
                        continue
                    n = int(n)
                if k in ("case_weight", "pallet_cases", "mfg_months") and n <= 0:
                    err(f"{label}は0より大きい数字にしてください（未設定にするなら空欄）")
                    continue
                values[k] = n
        out.append({"line": line, "code": code, "values": values})
    return out, errors
