import io
import json
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

os.environ["TANAOROSHI_DB"] = str(Path(tempfile.mkdtemp()) / "test.db")
os.environ["TANAOROSHI_ADMIN_PIN"] = "9999"

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import logic  # noqa: E402
from app.main import app  # noqa: E402

USER = {"X-User": "%E7%94%B0%E4%B8%AD"}  # 田中
ADMIN = {**USER, "X-Admin-Pin": "9999"}


# ---------- 判定ロジック ----------

def test_add_months_keeps_month_end():
    assert logic.add_months(date(2026, 9, 30), 1) == date(2026, 10, 31)
    assert logic.add_months(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert logic.add_months(date(2026, 1, 15), 6) == date(2026, 7, 15)
    assert logic.add_months(date(2026, 11, 10), 24) == date(2028, 11, 10)


@pytest.mark.parametrize("expiry, expected", [
    (date(2026, 10, 31), "ok"),       # 月末(9/30)からちょうど1ヶ月
    (date(2026, 10, 30), "warn"),     # 1ヶ月に1日足りない
    (date(2026, 9, 26), "warn"),
    (date(2026, 9, 24), "expired"),
    (None, "none"),
])
def test_judge(expiry, expected):
    assert logic.judge(expiry, "2026-09", date(2026, 9, 25)) == expected


def test_judge_months():
    assert logic.ok_limit("2026-10", 2) == date(2026, 12, 31)
    assert logic.judge(date(2026, 12, 30), "2026-10", date(2026, 10, 29), 2) == "warn"
    assert logic.judge(date(2026, 12, 31), "2026-10", date(2026, 10, 29), 2) == "ok"


def test_total_kg():
    assert logic.total_kg(10, 3, 4.5) == 34.5
    assert logic.total_kg(None, 3, 4.5) == 4.5


# ---------- API ----------

@pytest.fixture(scope="module")
def client():
    c = TestClient(app)
    assert c.post("/api/admin/staff", json={"name": "田中"}, headers=ADMIN).status_code == 200
    cat = c.post("/api/admin/categories", json={"grp": "植蛋", "name": "植蛋"}, headers=ADMIN).json()
    c.post("/api/admin/items", headers=ADMIN, json={"code": "04084", "name": "アペックス1000", "category_id": cat["id"], "case_weight": 7})
    c.post("/api/admin/items", headers=ADMIN, json={"code": "01009", "name": "上白糖", "category_id": cat["id"], "case_weight": 20,
                                                   "expiry_mode": "mfg", "mfg_months": 6})
    c.cat_id = cat["id"]
    return c


def items(client, ym="2026-09"):
    return {i["code"]: i for i in client.get(f"/api/months/{ym}/categories/{client.cat_id}").json()["items"]}


def test_requires_name_and_pin(client):
    assert client.put("/api/months/2026-09/items/1", json={}).status_code == 400
    assert client.get("/api/admin/items", headers=USER).status_code == 403


def test_save_entry_computes_total_and_status(client):
    it = items(client)["04084"]
    far = logic.add_months(date.today(), 3).isoformat()
    r = client.put(f"/api/months/2026-09/items/{it['id']}", headers=USER,
                   json={"room_cases": 2, "room_kg": 1.5, "wh_cases": 1, "wh_kg": 0, "expiry_kind": "賞", "expiry_date": far,
                         "action": "廃棄"})
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["total_kg"] == 22.5
    assert e["counted_by"] == "田中"
    assert e["counted_on"] == date.today().isoformat()  # 棚卸日は保存した日が自動で入る
    s = client.get("/api/months/2026-09/summary").json()["categories"][str(client.cat_id)]
    assert s["counted_from"] == s["counted_to"] == date.today().isoformat()
    assert e["action"] == ""  # 問題なければ対応は残さない


def test_loose_kg_parts_are_summed_and_kept(client):
    it = items(client)["04084"]
    far = logic.add_months(date.today(), 3).isoformat()
    r = client.put(f"/api/months/2026-09/items/{it['id']}", headers=USER,
                   json={"room_cases": 2, "room_kg_parts": [1.5, 2, 1], "wh_cases": 1, "wh_kg_parts": [0.2, 0, 0],
                         "expiry_kind": "賞", "expiry_date": far})
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["room_kg"] == 4.5 and e["wh_kg"] == 0.2
    assert e["total_kg"] == 25.7
    got = items(client)["04084"]
    assert json.loads(got["room_kg_parts"]) == [1.5, 2, 1]
    bad = client.put(f"/api/months/2026-09/items/{it['id']}", headers=USER,
                     json={"room_kg_parts": [1, 1, 1, 1], "expiry_kind": "賞", "expiry_date": far})
    assert bad.status_code == 422


def test_zero_stock_needs_no_expiry(client):
    for code in ("04084", "01009"):  # 期限を入力 / 製造日から計算
        it = items(client, "2026-06")[code]
        assert client.put(f"/api/months/2026-06/items/{it['id']}", headers=USER,
                          json={"room_cases": 1, "expiry_kind": "賞"}).status_code == 400
        r = client.put(f"/api/months/2026-06/items/{it['id']}", headers=USER, json={"room_cases": 0, "expiry_kind": "賞"})
        assert r.status_code == 200, r.text
        e = r.json()
        assert e["total_kg"] == 0 and e["status"] == "ok" and e["expiry_date"] is None


def test_pallets_count_as_cases(client):
    it = items(client, "2026-07")["04084"]
    far = logic.add_months(date.today(), 3).isoformat()
    body = {"room_cases": 1, "wh_pallets": 2, "wh_cases": 3, "wh_kg_parts": [0.5, 0, 0], "expiry_kind": "賞", "expiry_date": far}
    # 1パレットのケース数が未設定ならパレットは入れられない
    r = client.put(f"/api/months/2026-07/items/{it['id']}", headers=USER, json=body)
    assert r.status_code == 400
    master = {k: it[k] for k in ("code", "name", "category_id", "storage", "case_weight", "expiry_mode", "mfg_months", "note", "sort")}
    r = client.put(f"/api/admin/items/{it['id']}", headers=ADMIN, json={**master, "pallet_cases": 40})
    assert r.status_code == 200, r.text
    r = client.put(f"/api/months/2026-07/items/{it['id']}", headers=USER, json=body)
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["wh_pallets"] == 2 and e["pallet_cases"] == 40
    assert e["total_kg"] == (1 + 3 + 2 * 40) * 7 + 0.5
    ws = openpyxl.load_workbook(io.BytesIO(client.get("/api/months/2026-07/export.xlsx").content))["植蛋"]
    assert [ws.cell(row=6, column=c).value for c in range(5, 10)] == ["ケース", "端数kg", "パレット", "ケース", "端数kg"]
    row = [ws.cell(row=r, column=1).value for r in range(7, ws.max_row + 1)].index("04084") + 7
    assert [ws.cell(row=row, column=c).value for c in range(5, 10)] == [1, 0, 2, 3, 0.5]


def test_expired_entry_shows_in_alerts_and_action(client):
    ym = date.today().strftime("%Y-%m")
    it = items(client, ym)["04084"]
    past = date(2020, 1, 1).isoformat()
    r = client.put(f"/api/months/{ym}/items/{it['id']}", headers=USER,
                   json={"room_cases": 1, "expiry_kind": "凍", "expiry_date": past, "action": "廃棄", "action_note": "来週廃棄"})
    assert r.json()["status"] == "expired"
    alerts = client.get(f"/api/months/{ym}/alerts").json()["items"]
    assert [a["code"] for a in alerts] == ["04084"]
    assert alerts[0]["action"] == "廃棄"
    r = client.put(f"/api/months/{ym}/items/{it['id']}/action", headers=USER, json={"action": "移動", "action_note": "A倉庫"})
    assert r.json()["action"] == "移動"


def test_counted_on_can_be_corrected(client):
    it = items(client, "2026-08")["01009"]
    url = f"/api/months/2026-08/items/{it['id']}"
    body = {"room_cases": 1, "mfg_date": "2026-08-10"}
    assert client.put(url, headers=USER, json=body).json()["counted_on"] == date.today().isoformat()
    # あとから別の日に直せる
    assert client.put(url, headers=USER, json={**body, "counted_on": "2026-08-31"}).json()["counted_on"] == "2026-08-31"
    # 棚卸日を送らずに入れ直しても、直した日は変わらない
    assert client.put(url, headers=USER, json={**body, "room_cases": 2}).json()["counted_on"] == "2026-08-31"
    tomorrow = date.fromordinal(date.today().toordinal() + 1).isoformat()
    assert client.put(url, headers=USER, json={**body, "counted_on": tomorrow}).status_code == 400
    assert client.put(url, headers=USER, json={**body, "counted_on": "8月31日"}).status_code == 400
    s = client.get("/api/months/2026-08/summary").json()["categories"][str(client.cat_id)]
    assert s["counted_from"] == s["counted_to"] == "2026-08-31"
    assert "棚卸日2026-08-31" in client.get("/api/admin/history?entity=entry", headers=ADMIN).json()[0]["summary"]


def test_old_db_gets_counted_on(tmp_path, monkeypatch):
    import sqlite3
    from app import db
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE entries (id INTEGER PRIMARY KEY, ym TEXT, item_id INTEGER, counted_by TEXT, counted_at TEXT)")
    con.execute("INSERT INTO entries VALUES (1, '2026-09', 1, '田中', '2026-09-24T10:00:00')")
    con.commit()
    con.close()
    monkeypatch.setattr(db, "DB_PATH", path)
    db.init()
    with db.connect() as con:
        assert con.execute("SELECT counted_on FROM entries").fetchone()[0] == "2026-09-24"


def test_mfg_expiry(client):
    it = items(client)["01009"]
    assert client.put(f"/api/months/2026-09/items/{it['id']}", headers=USER, json={"room_cases": 1}).status_code == 400
    r = client.put(f"/api/months/2026-09/items/{it['id']}", headers=USER, json={"room_cases": 1, "mfg_date": "2026-08-10"})
    assert r.json()["expiry_date"] == "2027-02-10"


def test_expected_stock(client):
    it = items(client, "2026-10")["01009"]
    url = f"/api/months/2026-10/items/{it['id']}/expected"
    assert it["expected_kg"] is None
    # 棚卸の入力前でも入れられる
    r = client.put(url, headers=USER, json={"kg": 45.5}).json()
    assert r["kg"] == 45.5 and r["as_of"] == date.today().isoformat()  # 時点を省くと今日
    assert client.put(url, headers=USER, json={"kg": 40, "as_of": "2026-10-15"}).status_code == 200
    assert client.put(url, headers=USER, json={"kg": -1}).status_code == 422
    assert client.put(url, headers=USER, json={"kg": 40, "as_of": "10/15"}).status_code == 400
    it = items(client, "2026-10")["01009"]
    assert (it["expected_kg"], it["expected_as_of"]) == (40, "2026-10-15")
    assert items(client, "2026-11")["01009"]["expected_kg"] is None  # 月ごと
    hist = client.get("/api/admin/history?entity=expected", headers=ADMIN).json()
    assert [h["action"] for h in hist[:2]] == ["更新", "入力"]
    assert client.put(url, headers=USER, json={"kg": None}).status_code == 200
    assert items(client, "2026-10")["01009"]["expected_kg"] is None


def xlsx(rows) -> bytes:
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_expected_import_from_template(client):
    url = "/api/admin/months/2026-12/expected/import"
    r = client.get("/api/months/2026-12/expected.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[3]] == ["コード", "品名", "分類", "保管", "予想在庫(kg)", "時点"]
    row = next(r for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=1).value == "04084")
    ws.cell(row=row, column=5).value = 12.5
    ws.cell(row=row, column=6).value = datetime(2026, 11, 28)
    buf = io.BytesIO()
    ws.parent.save(buf)
    res = client.post(url, headers=ADMIN, content=buf.getvalue()).json()
    assert res == {"ok": True, "rows": 1, "changed": 1, "errors": []}
    it = items(client, "2026-12")["04084"]
    assert (it["expected_kg"], it["expected_as_of"]) == (12.5, "2026-11-28")
    assert client.get("/api/admin/history?entity=expected", headers=ADMIN).json()[0]["action"] == "取込"
    # 同じ値なら更新しない
    assert client.post(url, headers=ADMIN, content=buf.getvalue()).json()["changed"] == 0
    # 入っている時点がひな形に出る
    r = client.get("/api/months/2026-12/expected.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(row=row, column=6).value == datetime(2026, 11, 28)


def test_expected_template_code_order(client):
    r = client.get("/api/months/2026-12/expected.xlsx?order=code")
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    codes = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
    assert codes and codes == sorted(codes)


def test_expected_import_rules(client):
    url = "/api/admin/months/2026-12/expected/import"
    assert client.post(url, headers=USER, content=b"").status_code == 403
    # 数値になったコードも5桁に戻す。空欄は無視。「kg」付きの文字も読む
    res = client.post(url, headers=ADMIN, content=xlsx([["メモ"], ["コード", "予想在庫"], [1009, "30kg"], ["04084", None]])).json()
    assert res["ok"] and res["rows"] == 1
    assert items(client, "2026-12")["01009"]["expected_kg"] == 30
    assert items(client, "2026-12")["01009"]["expected_as_of"] == date.today().isoformat()
    # 時点が空欄の行は画面で選んだ日。時点だけ変わっても更新する
    res = client.post(url + "?as_of=2026-11-30", headers=ADMIN, content=xlsx([["コード", "予想在庫", "時点"], ["01009", 30, None], ["04084", 5, "2026/11/20"]])).json()
    assert res["ok"] and res["changed"] == 2
    got = items(client, "2026-12")
    assert got["01009"]["expected_as_of"] == "2026-11-30" and got["04084"]["expected_as_of"] == "2026-11-20"
    res = client.post(url, headers=ADMIN, content=xlsx([["コード", "予想在庫", "時点"], ["01009", 30, "そのうち"]])).json()
    assert not res["ok"] and "時点" in res["errors"][0]
    assert client.post(url + "?as_of=bad", headers=ADMIN, content=b"").status_code == 400
    # エラーが1件でもあれば何も書かない
    res = client.post(url, headers=ADMIN, content=xlsx([["コード", "予想在庫(kg)"], ["01009", 99], ["99999", 1], ["04084", "たくさん"]])).json()
    assert not res["ok"] and len(res["errors"]) == 2
    assert items(client, "2026-12")["01009"]["expected_kg"] == 30
    assert not client.post(url, headers=ADMIN, content=b"not excel").json()["ok"]
    assert "見出し" in client.post(url, headers=ADMIN, content=xlsx([["a", "b"]])).json()["errors"][0]


def test_item_change_and_discontinue_are_logged(client):
    it = items(client)["04084"]
    body = {"code": "04084", "name": "アペックス1000", "category_id": client.cat_id, "case_weight": 8, "active": False}
    assert client.put(f"/api/admin/items/{it['id']}", headers=ADMIN, json=body).status_code == 200
    hist = client.get("/api/admin/history?entity=item", headers=ADMIN).json()
    assert hist[0]["action"] == "終売"
    assert "ケース重量: 7 → 8" in hist[0]["summary"]
    # 終売でも、入力済みの月には残る。未入力の月には出ない
    assert "04084" in items(client, "2026-09")
    assert "04084" not in items(client, "2030-01")


def test_approval_flow(client):
    url = f"/api/months/2026-09/categories/{client.cat_id}/approvals"
    assert client.post(url, headers=USER, json={"role": "担当者"}).status_code == 200
    assert client.post(url, headers=USER, json={"role": "担当者"}).status_code == 409
    assert client.post(url, headers=USER, json={"role": "社長"}).status_code == 400
    s = client.get("/api/months/2026-09/summary").json()["categories"][str(client.cat_id)]
    assert s["approvals"] == {"担当者": "田中"}
    assert client.delete(f"{url}/%E6%8B%85%E5%BD%93%E8%80%85", headers=USER).status_code == 200


def test_static_is_not_cached(client):
    assert client.get("/static/app.js").headers["cache-control"] == "no-cache"


def test_export_xlsx(client):
    r = client.get("/api/months/2026-09/export.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["期限警告一覧", "植蛋"]
    ws = wb["植蛋"]
    codes = [ws.cell(row=r, column=1).value for r in range(7, ws.max_row + 1)]
    assert set(codes) == {"04084", "01009"}
    today = date.today()
    assert f"棚卸日：{today.year}/{today.month}/{today.day}" in ws["A2"].value
    heads = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
    col = heads.index("棚卸日") + 1
    row = codes.index("04084") + 7
    assert ws.cell(row=row, column=col).value == f"{today.year}/{today.month}/{today.day}"


def test_items_excel_round_trip(client):
    assert client.get("/api/admin/items.xlsx", headers=USER).status_code == 403
    r = client.get("/api/admin/items.xlsx", headers=ADMIN)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["品目"]
    heads = [c.value for c in ws[3]]
    col = {h: i + 1 for i, h in enumerate(heads)}
    rows = {ws.cell(row=i, column=1).value: i for i in range(4, ws.max_row + 1)}
    assert set(rows) == {"04084", "01009"}
    assert ws.cell(row=rows["01009"], column=col["期限の種類"]).value == "製造日から計算"
    url = "/api/admin/items/import"
    # そのまま戻しても何も変わらない
    buf = io.BytesIO(); wb.save(buf)
    r = client.post(url, headers=ADMIN, content=buf.getvalue()).json()
    assert r["ok"] and r["changes"] == []

    ws.cell(row=rows["01009"], column=col["1パレットのケース数"], value=50)
    ws.cell(row=rows["01009"], column=col["ケース重量(kg)"]).value = None  # 空欄=未設定
    assert ws.cell(row=rows["04084"], column=col["状態"]).value == "終売"  # 前のテストで終売にした
    ws.cell(row=rows["04084"], column=col["状態"], value="有効")
    new = ws.max_row + 1
    for h, v in {"コード": 2001, "品名": "新しい粉", "分類": "植蛋 / 植蛋", "保管": "冷蔵", "期限の種類": "期限指定無し"}.items():
        ws.cell(row=new, column=col[h], value=v)
    buf = io.BytesIO(); wb.save(buf)
    data = buf.getvalue()
    r = client.post(url, headers=ADMIN, content=data).json()  # 確認だけ
    assert r["ok"] and (r["added"], r["changed"]) == (1, 2)
    assert {c["code"]: c["kind"] for c in r["changes"]} == {"01009": "変更", "04084": "変更", "02001": "追加"}
    assert items(client, "2026-06")["01009"]["case_weight"] == 20  # まだ反映しない
    r = client.post(url + "?apply=1", headers=ADMIN, content=data).json()
    assert r["ok"] and r["applied"]
    master = {i["code"]: i for i in client.get("/api/admin/items", headers=ADMIN).json()}
    assert master["01009"]["pallet_cases"] == 50 and master["01009"]["case_weight"] is None
    assert master["04084"]["active"] == 1
    assert master["02001"]["storage"] == "冷蔵" and master["02001"]["expiry_mode"] == "none"
    hist = client.get("/api/admin/history?entity=item", headers=ADMIN).json()
    assert {"取込追加", "取込変更"} <= {h["action"] for h in hist}

    # エラーが1件でもあれば何も変えない
    bad = xlsx([["コード", "品名", "保管", "ケース重量(kg)"], ["01009", "上白糖", "冷蔵", 5], ["01009", "", "常温", None],
                ["9999", "分類なし", "常温", 1], ["02001", "", "地下", "重い"]])
    r = client.post(url + "?apply=1", headers=ADMIN, content=bad).json()
    assert not r["ok"] and len(r["errors"]) == 4
    master2 = {i["code"]: i for i in client.get("/api/admin/items", headers=ADMIN).json()}
    assert master2["01009"] == master["01009"]



def test_items_import_circled_digit_category(client):
    """「粉体①」のような分類名・品名も、出力したExcelをそのまま戻せる。"""
    cat = client.post("/api/admin/categories", json={"grp": "粉体資材", "name": "粉体①"}, headers=ADMIN).json()
    url = "/api/admin/items/import"
    data = xlsx([["コード", "品名", "分類"], ["03001", "粉①　特製", "粉体資材 / 粉体①"]])
    r = client.post(url + "?apply=1", headers=ADMIN, content=data).json()
    assert r["ok"], r["errors"]
    master = {i["code"]: i for i in client.get("/api/admin/items", headers=ADMIN).json()}
    assert master["03001"]["category_id"] == cat["id"] and master["03001"]["name"] == "粉①　特製"
    # Excelで出力したものを戻しても変更なし
    out = client.get("/api/admin/items.xlsx", headers=ADMIN).content
    r = client.post(url, headers=ADMIN, content=out).json()
    assert r["ok"] and r["changes"] == [], r

# ---------- バックアップ ----------

def test_backup_copies_wal_and_prunes_old(tmp_path):
    import sqlite3
    from tools import backup

    src = tmp_path / "src.db"
    con = sqlite3.connect(src)
    con.execute("PRAGMA journal_mode = WAL")
    con.execute("CREATE TABLE t (v)")
    con.execute("INSERT INTO t VALUES ('最新')")
    con.commit()  # 接続を開いたままなので、この行は WAL に残っている

    dest = tmp_path / "bk"
    out = backup.backup(src, dest, date(2026, 10, 9))
    con.close()
    assert out.name == "tanaoroshi_20261009.db"
    b = sqlite3.connect(out)
    assert b.execute("SELECT v FROM t").fetchall() == [("最新",)]
    b.close()
    backup.backup(src, dest, date(2026, 10, 9))  # 同じ日は上書き

    for name in ["tanaoroshi_20250908.db", "tanaoroshi_20250909.db", "tanaoroshi_20251010.db", "memo.db"]:
        (dest / name).write_bytes(b"")
    removed = backup.prune(dest, date(2026, 10, 9))
    assert [p.name for p in removed] == ["tanaoroshi_20250908.db"]  # 13ヶ月前の 2025/9/9 より前だけ消す
    assert sorted(p.name for p in dest.iterdir()) == [
        "memo.db", "tanaoroshi_20250909.db", "tanaoroshi_20251010.db", "tanaoroshi_20261009.db"]


def test_item_docs(client):
    it = items(client)["01009"]
    url = f"/api/admin/items/{it['id']}/docs?name=" + "%E8%A6%8F%E6%A0%BC%E6%9B%B8.pdf"  # 規格書.pdf
    pdf = b"%PDF-1.4\n%test\n"
    assert client.post(url, headers=USER, content=pdf).status_code == 403
    assert client.post(url, headers=ADMIN, content=b"not a pdf").status_code == 400
    doc = client.post(url, headers=ADMIN, content=pdf).json()
    assert doc["name"] == "規格書.pdf" and doc["size"] == len(pdf)
    # 誰でも一覧を見て開ける。入力画面の品目にも件数が出る
    assert [d["id"] for d in client.get(f"/api/items/{it['id']}/docs").json()] == [doc["id"]]
    r = client.get(f"/api/docs/{doc['id']}")
    assert r.content == pdf and r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"].startswith("inline")
    assert items(client)["01009"]["doc_count"] == 1
    assert client.get("/api/admin/history?entity=doc", headers=ADMIN).json()[0]["action"] == "追加"
    # 削除するとファイルも消える
    assert client.delete(f"/api/admin/docs/{doc['id']}", headers=ADMIN).json()["ok"]
    assert client.get(f"/api/docs/{doc['id']}").status_code == 404
    assert items(client)["01009"]["doc_count"] == 0


def test_backup_docs_copies_only_new(tmp_path):
    from tools import backup
    src = tmp_path / "docs"
    src.mkdir()
    (src / "1.pdf").write_bytes(b"%PDF-1")
    dest = tmp_path / "bk"
    assert backup.backup_docs(src, dest) == 1
    (src / "2.pdf").write_bytes(b"%PDF-2")
    assert backup.backup_docs(src, dest) == 1
    assert sorted(p.name for p in (dest / "docs").iterdir()) == ["1.pdf", "2.pdf"]
    assert backup.backup_docs(tmp_path / "none", dest) == 0


def minimal_pdf(pages: int) -> bytes:
    """白紙のPDFを作る（xref の位置も正しく書く）。"""
    kids = " ".join(f"{3 + i} 0 R" for i in range(pages))
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", f"<< /Type /Pages /Kids [{kids}] /Count {pages} >>"]
    objs += ["<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] >>"] * pages
    out, offsets = b"%PDF-1.4\n", []
    for n, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{x:010d} 00000 n \n" for x in offsets).encode()
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


def test_png_writer():
    import zlib
    from app import pdf_pages
    data = pdf_pages.png(2, 1, 3, 8, bytes([255, 0, 0, 0, 255, 0, 9, 9]))  # stride の余り(9,9)は捨てる
    assert data.startswith(b"\x89PNG\r\n\x1a\n") and data[12:16] == b"IHDR"
    idat = data.index(b"IDAT")
    size = int.from_bytes(data[idat - 4:idat], "big")
    assert zlib.decompress(data[idat + 4:idat + 4 + size]) == bytes([0, 255, 0, 0, 0, 255, 0])


def test_doc_page_images(client):
    pytest.importorskip("pypdfium2")
    it = items(client)["04084"]
    doc = client.post(f"/api/admin/items/{it['id']}/docs?name=a.pdf", headers=ADMIN, content=minimal_pdf(2)).json()
    assert client.get(f"/api/items/{it['id']}/docs").json()[0]["pages"] == 2
    thumb = client.get(f"/api/docs/{doc['id']}/pages/1.png")
    large = client.get(f"/api/docs/{doc['id']}/pages/2.png?size=large")
    assert thumb.headers["content-type"] == "image/png" and int.from_bytes(thumb.content[16:20], "big") == 240
    assert int.from_bytes(large.content[16:20], "big") == 1200
    assert client.get(f"/api/docs/{doc['id']}/pages/3.png").status_code == 404
    client.delete(f"/api/admin/docs/{doc['id']}", headers=ADMIN)


def test_category_warn_months_rejudges_this_month(client):
    cat = client.post("/api/admin/categories", json={"grp": "原料肉", "name": "原料肉"}, headers=ADMIN).json()
    assert cat["warn_months"] == 1
    it = client.post("/api/admin/items", headers=ADMIN, json={"code": "M001", "name": "豚肉", "category_id": cat["id"],
                                                              "case_weight": 10}).json()
    ym = date.today().strftime("%Y-%m")
    expiry = logic.add_months(logic.month_end(ym), 1).isoformat()  # 1ヶ月ならレ、2ヶ月なら×
    url = f"/api/months/{ym}/items/{it['id']}"
    client.put(url, headers=USER, json={"room_cases": 1, "expiry_date": expiry})
    old = client.put(f"/api/months/2020-01/items/{it['id']}", headers=USER,
                     json={"room_cases": 1, "expiry_date": "2020-02-29"}).json()

    body = {"grp": "原料肉", "name": "原料肉", "warn_months": 2}
    assert client.put(f"/api/admin/categories/{cat['id']}", headers=ADMIN, json={**body, "warn_months": 0}).status_code == 422
    assert client.put(f"/api/admin/categories/{cat['id']}", headers=ADMIN, json=body).json()["warn_months"] == 2
    d = client.get(f"/api/months/{ym}/categories/{cat['id']}").json()
    assert d["ok_limit"] == logic.ok_limit(ym, 2).isoformat()
    assert d["items"][0]["status"] == "warn" and d["items"][0]["warn_months"] == 2
    # 過去の月は判定し直さない
    assert client.get(f"/api/months/2020-01/categories/{cat['id']}").json()["items"][0]["status"] == old["status"]
    assert "再判定" in [h["action"] for h in client.get("/api/admin/history?entity=entry", headers=ADMIN).json()]
    # 新しく保存するときも分類の月数で判定する
    assert client.put(url, headers=USER, json={"room_cases": 2, "expiry_date": expiry}).json()["status"] == "warn"
